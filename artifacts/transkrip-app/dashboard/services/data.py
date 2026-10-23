"""Pembaca data hasil scraping untuk dashboard Django (port dari webapp/data.py).

Semua path kini mengikuti ecampuz.config — data berada di project transkrip-app
(pada <BASE_DIR>/data/ecampuz, atau di-override via env ECAMPUZ_ROOT).
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from ecampuz.config import DATA_DIR, OUT_DIR, STATE_DIR

ROOT = DATA_DIR
DATA_DIR = DATA_DIR  # alias biar pola pemakaian lama tetap kompatibel
OUT_DIR = OUT_DIR
STATE_DIR = STATE_DIR

# Default path template Excel — relatif terhadap data/ecampuz.
DEFAULT_FILLED_EXCEL = DATA_DIR / "Template_Transkrip_2026_FIXED (5) - FILLED.xlsx"
DEFAULT_ORIGINAL_EXCEL = DATA_DIR / "Template_Transkrip_2026_FIXED (5).xlsx"
WEB_CONFIG_PATH = STATE_DIR / "web.json"

TEMPLATE_SHEET = "Nilai MHS 2026"


def web_config() -> dict:
    default = {
        "filled_excel": str(DEFAULT_FILLED_EXCEL),
        "nim_file": "nim_2024_sheet.txt",
        "angkatan": "2024",
        "prodi": "all",
    }
    if not WEB_CONFIG_PATH.exists():
        return default
    try:
        data = json.loads(WEB_CONFIG_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return default
    merged = {**default, **data}
    return merged


def save_web_config(cfg: dict) -> None:
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    WEB_CONFIG_PATH.write_text(
        json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8"
    )


def _read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    try:
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    except (OSError, csv.Error):
        return []


def resolve_nim_file(path_value) -> Path:
    """Resolve path file NIM ke abs; toleransi path lama 'data/nim_...'."""
    p = Path(str(path_value or ""))
    if p.is_absolute() and p.exists():
        return p
    candidate = ROOT / p
    if candidate.exists():
        return candidate
    # path lama (relatif ke root scraper lama) — cari berdasarkan nama file
    if p.name:
        fallback = ROOT / p.name
        if fallback.exists():
            return fallback
    return candidate


def dhs_rows() -> list[dict]:
    return _read_csv(OUT_DIR / "dhs_semua.csv")


def prodi_options() -> list[dict[str, str]]:
    """Daftar prodi resmi yang diterima oleh eCampuz, termasuk opsi valid yang bisa dipilih."""
    choices = [
        {"value": "all", "label": "Semua Program Studi"},
        {"value": "MTg%3D", "label": "ILMU KEPERAWATAN - S1 Reguler"},
        {"value": "NDg%3D", "label": "INFORMATIKA - S1 Reguler"},
        {"value": "Mzg%3D", "label": "KEBIDANAN - D3 Reguler"},
        {"value": "Njg%3D", "label": "MANAJEMEN RITEL - S1 Reguler"},
        {"value": "Nzg%3D", "label": "PERPAJAKAN - S1 Reguler"},
        {"value": "Mjg%3D", "label": "PROFESI NERS - Profesi Reguler"},
        {"value": "NTg%3D", "label": "SISTEM INFORMASI - S1 Reguler"},
    ]

    seen: set[str] = set()
    out: list[dict[str, str]] = []
    for item in choices:
        value = item["value"]
        if value in seen:
            continue
        seen.add(value)
        out.append(item)
    return out


def student_summaries() -> dict[str, dict]:
    out: dict[str, dict] = {}
    for r in dhs_rows():
        nim = r.get("nim", "").strip()
        if not nim:
            continue
        s = out.setdefault(
            nim,
            {
                "nim": nim,
                "nama": r.get("nama", ""),
                "prodi": r.get("prodi", ""),
                "angkatan": r.get("angkatan", ""),
                "sks": 0.0,
                "mutu": 0.0,
                "jumlah_mk": 0,
            },
        )
        if r.get("nama"):
            s["nama"] = r["nama"]
        if r.get("prodi"):
            s["prodi"] = r["prodi"]
        if r.get("angkatan"):
            s["angkatan"] = r["angkatan"]
        try:
            sks = float(r["sks"])
            bobot = float(r["bobot"])
        except (KeyError, TypeError, ValueError):
            continue
        s["sks"] += sks
        s["mutu"] += sks * bobot
        s["jumlah_mk"] += 1
    for s in out.values():
        s["ipk"] = round(s["mutu"] / s["sks"], 3) if s["sks"] else 0.0
    return out


def nim_transcript(nim: str) -> dict | None:
    path = OUT_DIR / f"dhs_{nim}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None


def mapping_rows() -> list[dict]:
    raw = _read_csv(DATA_DIR / "mapping_draft.csv")
    if not raw:
        return []

    # Lookup nama mata kuliah kurikulum dari database untuk melengkapi nama template yang kosong
    db_mks: dict[str, str] = {}
    try:
        from core.models import MataKuliahKurikulum
        db_mks = {
            m.kode_mk.strip().upper(): m.nama_mk.strip()
            for m in MataKuliahKurikulum.objects.all()
        }
    except Exception:
        db_mks = {}

    out = []
    for r in raw:
        kode_dhs = (r.get("eA_kode") or r.get("kode_dhs") or "").strip()
        nama_dhs = (r.get("eA_nama") or r.get("nama_dhs") or "").strip()
        sks = (r.get("eA_sks") or r.get("sks") or "").strip()
        n_nim = (r.get("n_nim") or r.get("jumlah_mhs") or "").strip()
        kode_tpl = (r.get("P24_kode") or r.get("kode_template") or "").strip()
        nama_tpl = (r.get("P24_nama") or r.get("nama_template") or "").strip()
        status = (r.get("status") or "").strip().upper()
        kandidat = (r.get("kandidat") or r.get("rekomendasi") or "").strip()

        # Jika nama template kosong di CSV (misal pada baris APPROVED-REC), ambil dari DB master kurikulum
        if not nama_tpl and kode_tpl:
            nama_tpl = db_mks.get(kode_tpl.upper(), "")

        # Buat keterangan kesimpulan / catatan rekomendasi
        if kandidat:
            kandidat_nama = db_mks.get(kandidat.upper(), "")
            rekomendasi = f"{kandidat} ({kandidat_nama})" if kandidat_nama else kandidat
        elif status == "APPROVED-EXACT":
            rekomendasi = "Nama cocok persis"
        elif status == "APPROVED-REC":
            rekomendasi = f"Dipetakan ke {kode_tpl}" if kode_tpl else "Rekomendasi padanan"
        elif status == "SKIP":
            rekomendasi = "Dilewati (tidak ada padanan langsung / ambigu)"
        else:
            rekomendasi = "-"

        out.append({
            # Key normalized untuk kemudahan rendering template
            "kode_dhs": kode_dhs,
            "nama_dhs": nama_dhs,
            "sks_dhs": sks,
            "n_nim": n_nim,
            "kode_template": kode_tpl,
            "nama_template": nama_tpl,
            "rekomendasi": rekomendasi,
            "status": status,
            # Backward-compatibility alias
            "eA_kode": kode_dhs,
            "eA_nama": nama_dhs,
            "eA_sks": sks,
            "P24_kode": kode_tpl,
            "P24_nama": nama_tpl,
            "kandidat": kandidat,
            "kolom": sks,
        })
    return out


def checkpoint(angkatan: str) -> dict:
    path = OUT_DIR / f"dhs_state_{angkatan}.json"
    if not path.exists():
        return {"done": [], "failed": {}}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {"done": [], "failed": {}}
    return {"done": list(data.get("done", [])), "failed": dict(data.get("failed", {}))}


_xl_cache: dict = {"mtime": -1, "data": None}


def _excel_record(row: tuple) -> dict:
    def v(i):
        # read_only=True -> openpyxl memberi ReadOnlyCell, bukan nilai biasa.
        cell = row[i]
        val = getattr(cell, "value", cell)
        return "" if val is None else val

    rec = {
        "no": v(0),
        "nim": str(v(1)).strip(),
        "nama": v(2),
        "kode": str(v(3)).strip() if v(3) != "" else "",
        "nama_id": v(4),
        "nama_en": v(5),
        "sks": v(6),
        "semester": "" if v(7) == "" else str(v(7)),
        "indeks": v(8),
        "huruf": str(v(9)).strip() if v(9) != "" else "",
        "sks_indeks": v(10),
    }
    rec["terisi"] = bool(rec["semester"] and rec["huruf"] and rec["indeks"] != "")
    return rec


def excel_data_path() -> Path:
    cfg = web_config()
    path = Path(cfg.get("filled_excel") or "") if cfg.get("filled_excel") else None
    if path and path.exists():
        return path
    if DEFAULT_FILLED_EXCEL.exists():
        return DEFAULT_FILLED_EXCEL
    return DEFAULT_ORIGINAL_EXCEL


def excel_data() -> dict:
    path = excel_data_path()
    if not path.exists():
        return {}
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return {}
    if _xl_cache["data"] is not None and _xl_cache["mtime"] == mtime:
        return _xl_cache["data"]

    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if TEMPLATE_SHEET not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[TEMPLATE_SHEET]
    by_nim: dict[str, list] = {}
    for row in ws.iter_rows(min_row=2):
        rec = _excel_record(row)
        if not rec["nim"].isdigit():
            continue
        by_nim.setdefault(rec["nim"], []).append(rec)
    wb.close()
    _xl_cache["mtime"] = mtime
    _xl_cache["data"] = by_nim
    return by_nim


def excel_filled_stats() -> dict:
    total = 0
    filled = 0
    for rows in excel_data().values():
        total += len(rows)
        filled += sum(1 for r in rows if r["terisi"])
    return {"total": total, "terisi": filled, "kosong": total - filled}


def audit_tail(limit: int = 300) -> list[dict]:
    path = STATE_DIR / "audit.jsonl"
    if not path.exists():
        return []
    try:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []
    out = []
    for line in lines[-limit:]:
        try:
            out.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return out


def session_file_exists() -> bool:
    return (STATE_DIR / "session.json").exists()