"""Template creation and preflight validation for eCampuz graduation imports."""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

IMPORT_SHEET_NAME = "Import data kelulusan"
IMPORT_HEADER_ROW = 6
IMPORT_HEADERS = [
    "No",
    "NIM",
    "Nama",
    "Status Lulus",
    "Semester Lulus",
    "Tanggal Yudisium",
    "No Surat Yudisium",
    "Tanggal Keluar/Lulus",
    "Jalur Lulus (Skripsi/Non Skripsi)",
    "Nomor Ijazah",
    "Tanggal Ijazah",
    "No ijazah lain",
]
_ALLOWED_STATUS = {"lulus", "belum lulus"}
_ALLOWED_PATHS = {"skripsi", "non skripsi"}
_DATE_FORMATS = ("%d-%m-%Y",)
_IJAZAH_RE = re.compile(r"^[0-9A-Za-z][0-9A-Za-z./-]{2,99}$")


def build_graduation_template() -> bytes:
    """Return a clean BIFF .xls workbook matching the campus import layout."""
    import xlwt

    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet(IMPORT_SHEET_NAME)

    title_style = xlwt.easyxf(
        "font: bold on, height 260; align: horiz center, vert center;"
        "pattern: pattern solid, fore_colour ice_blue;"
        "borders: bottom thin, bottom_colour gray25"
    )
    note_style = xlwt.easyxf("font: bold on, colour dark_red; align: vert center")
    guide_style = xlwt.easyxf("font: colour dark_blue; align: vert center, wrap on")
    guide_example_style = xlwt.easyxf(
        "font: colour dark_green; align: vert center, wrap on"
    )
    header_style = xlwt.easyxf(
        "font: bold on, colour white; align: horiz center, vert center, wrap on;"
        "pattern: pattern solid, fore_colour dark_blue;"
        "borders: left thin, right thin, top thin, bottom thin,"
        "left_colour gray25, right_colour gray25, top_colour gray25, bottom_colour gray25"
    )
    text_style = xlwt.easyxf(
        "align: vert center; borders: bottom thin, bottom_colour gray25"
    )
    nim_style = xlwt.easyxf(
        "align: vert center; borders: bottom thin, bottom_colour gray25",
        num_format_str="@",
    )
    number_style = xlwt.easyxf(
        "align: horiz center, vert center; borders: bottom thin, bottom_colour gray25",
        num_format_str="0",
    )
    # eCampuz rejects real Excel date cells during kelulusan import. Keep the
    # template date columns as text so typed values are sent as dd-mm-yyyy.
    date_text_style = xlwt.easyxf(
        "align: vert center; borders: bottom thin, bottom_colour gray25",
        num_format_str="@",
    )

    sheet.write_merge(0, 0, 0, 11, "DAFTAR MAHASISWA LULUS", title_style)
    sheet.write(1, 0, "Note", note_style)
    sheet.write_merge(
        1,
        1,
        1,
        11,
        "Isi satu baris untuk setiap mahasiswa. Jangan mengubah nama sheet, posisi header, atau urutan kolom.",
        guide_style,
    )
    sheet.write(2, 3, "Status Lulus", guide_style)
    sheet.write(2, 4, "Semester Lulus", guide_style)
    sheet.write(2, 5, "Tanggal Yudisium", guide_style)
    sheet.write(2, 6, "No Surat Yudisium", guide_style)
    sheet.write(2, 7, "Tanggal Keluar/Lulus", guide_style)
    sheet.write(2, 8, "Jalur Lulus", guide_style)
    sheet.write(2, 9, "Nomor Ijazah", guide_style)
    sheet.write(2, 10, "Tanggal Ijazah", guide_style)
    sheet.write(2, 11, "No Ijazah Lain", guide_style)
    for column, value in {
        3: "Lulus / Belum Lulus",
        4: "Contoh: Ganjil 2026-2027",
        5: "Teks dd-mm-yyyy",
        6: "Nomor surat / teks",
        7: "Teks dd-mm-yyyy",
        8: "Skripsi / Non Skripsi",
        9: "142012023000112",
        10: "Teks dd-mm-yyyy",
        11: "Opsional, teks",
    }.items():
        sheet.write(3, column, value, guide_example_style)
    sheet.write(4, 3, "Isi status persis sesuai pilihan.", guide_example_style)
    sheet.write(4, 4, "Contoh format semester.", guide_example_style)
    sheet.write(4, 5, "Teks, contoh: 10-03-2027", guide_example_style)
    sheet.write(4, 7, "Teks, contoh: 10-03-2027", guide_example_style)
    sheet.write(4, 8, "Pilih salah satu jalur.", guide_example_style)
    sheet.write(4, 9, "Nomor ijazah PDDIKTI/Neo Feeder (bukan A/B).", guide_example_style)
    sheet.write(4, 10, "Teks, contoh: 10-03-2027", guide_example_style)
    sheet.write_merge(
        4,
        4,
        0,
        2,
        "Baris data dimulai pada baris 8. NIM disimpan sebagai teks.",
        guide_example_style,
    )

    sheet.row(0).height = 792
    for row_index in range(1, IMPORT_HEADER_ROW):
        sheet.row(row_index).height = 420
    sheet.row(IMPORT_HEADER_ROW).height = 720
    for column, header in enumerate(IMPORT_HEADERS):
        sheet.write(IMPORT_HEADER_ROW, column, header, header_style)

    widths = [8, 16, 28, 18, 23, 19, 24, 22, 31, 18, 18, 24]
    for column, width in enumerate(widths):
        sheet.col(column).width = width * 256

    for row_index in range(IMPORT_HEADER_ROW + 1, IMPORT_HEADER_ROW + 101):
        sheet.write(row_index, 0, "", number_style)
        sheet.write(row_index, 1, "", nim_style)
        sheet.write(row_index, 2, "", text_style)
        for column in (3, 4, 6, 8, 9, 11):
            sheet.write(row_index, column, "", text_style)
        for column in (5, 7, 10):
            sheet.write(row_index, column, "", date_text_style)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_roster_template() -> bytes:
    """Return a clean .xlsx workbook with student roster headers and example data."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster Yudisium"

    headers = [
        "No",
        "NIM",
        "Nama Mahasiswa",
        "Program Studi",
        "Angkatan",
        "Keterangan",
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[1].height = 26

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    sample_data = [
        (1, "2021010001", "Ahmad Fadillah", "Informatika", 2021, "Siap Daftar"),
        (2, "2021010002", "Siti Nurhaliza", "Sistem Informasi", 2021, "Siap Daftar"),
        (3, "2021010003", "Budi Santoso", "Informatika", 2021, "Siap Daftar"),
    ]

    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx == 2:
                cell.number_format = "@"
                cell.value = str(value)
                cell.alignment = center_align
            elif col_idx in (1, 5):
                cell.value = value
                cell.alignment = center_align
            else:
                cell.value = str(value)
                cell.alignment = left_align

    col_widths = {
        1: 8,
        2: 18,
        3: 28,
        4: 24,
        5: 14,
        6: 22,
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _display_value(value: Any, cell_type: int, datemode: int) -> str:
    import xlrd

    if value is None or value == "":
        return ""
    if cell_type == xlrd.XL_CELL_DATE:
        parsed = xlrd.xldate_as_datetime(value, datemode)
        return parsed.strftime("%d-%m-%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _valid_date(value: Any, cell_type: int, datemode: int) -> bool:
    import xlrd

    if value in (None, ""):
        return True
    # eCampuz menolak sel tanggal Excel pada import kelulusan; template
    # memakai sel teks dan validator harus menolak format sel tanggal juga.
    if cell_type == xlrd.XL_CELL_DATE:
        return False
    text = str(value).strip()
    if not re.fullmatch(r"\d{2}-\d{2}-\d{4}", text):
        return False
    for fmt in _DATE_FORMATS:
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def _valid_ijazah_number(value: Any) -> bool:
    text = str(value or "").strip()
    if not text or text.casefold() in {"a", "b"}:
        return False
    if not any(ch.isdigit() for ch in text):
        return False
    return bool(_IJAZAH_RE.fullmatch(text))


def validate_graduation_import(
    content: bytes,
    filename: str = "graduation.xls",
) -> dict[str, Any]:
    """Validate a completed legacy .xls without exposing student names."""
    import xlrd

    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    if not filename.casefold().endswith(".xls"):
        return {
            "valid": False,
            "total": 0,
            "errors": [{"row": None, "field": "file", "message": "File import eCampuz harus berformat .xls."}],
            "warnings": [],
        }
    try:
        workbook = xlrd.open_workbook(file_contents=content, formatting_info=False)
    except (xlrd.XLRDError, ValueError, OSError) as exc:
        return {
            "valid": False,
            "total": 0,
            "errors": [{"row": None, "field": "file", "message": f"File bukan workbook XLS yang bisa dibaca: {exc}"}],
            "warnings": [],
        }

    if IMPORT_SHEET_NAME not in workbook.sheet_names():
        workbook.release_resources()
        return {
            "valid": False,
            "total": 0,
            "errors": [{"row": None, "field": "sheet", "message": f"Nama sheet harus tepat {IMPORT_SHEET_NAME!r}."}],
            "warnings": [],
        }

    sheet = workbook.sheet_by_name(IMPORT_SHEET_NAME)
    if sheet.nrows <= IMPORT_HEADER_ROW:
        workbook.release_resources()
        return {
            "valid": False,
            "total": 0,
            "errors": [{"row": None, "field": "header", "message": "Header baris 7 atau baris data tidak ditemukan."}],
            "warnings": [],
        }
    actual_headers = [
        _display_value(sheet.cell_value(IMPORT_HEADER_ROW, column), sheet.cell_type(IMPORT_HEADER_ROW, column), workbook.datemode)
        for column in range(len(IMPORT_HEADERS))
    ]
    if actual_headers != IMPORT_HEADERS:
        workbook.release_resources()
        return {
            "valid": False,
            "total": 0,
            "errors": [{
                "row": IMPORT_HEADER_ROW + 1,
                "field": "header",
                "message": "Header baris 7 tidak sama dengan format eCampuz. Unduh template resmi dari halaman ini.",
            }],
            "warnings": [],
        }

    seen_nims: set[str] = set()
    records: list[dict[str, Any]] = []
    for row_index in range(IMPORT_HEADER_ROW + 1, sheet.nrows):
        raw = [sheet.cell_value(row_index, col) for col in range(len(IMPORT_HEADERS))]
        types = [sheet.cell_type(row_index, col) for col in range(len(IMPORT_HEADERS))]
        row = [_display_value(value, types[index], workbook.datemode) for index, value in enumerate(raw)]
        if not any(value for value in row):
            continue
        excel_row = row_index + 1
        record_errors: list[dict[str, Any]] = []

        number_text = row[0]
        try:
            number = int(number_text)
            if number < 1:
                raise ValueError
            if number != len(records) + 1:
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[0], "message": f"Nomor baris harus urut; nilai berikutnya {len(records) + 1}."})
        except ValueError:
            record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[0], "message": "Nomor harus bilangan bulat positif."})
        nim = row[1]
        if not re.fullmatch(r"[0-9]{6,20}", nim):
            record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[1], "message": "NIM wajib berisi 6-20 digit angka; gunakan sel teks agar nol di depan tidak hilang."})
        elif nim in seen_nims:
            record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[1], "message": "NIM duplikat di dalam file."})
        else:
            seen_nims.add(nim)
        if not row[2]:
            record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[2], "message": "Nama wajib diisi."})

        status = row[3].casefold()
        if status not in _ALLOWED_STATUS:
            record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[3], "message": "Isi tepat 'Lulus' atau 'Belum Lulus'."})
        if status == "lulus":
            semester = row[4]
            if not semester or not re.search(r"\d{4}\s*[-/]\s*\d{4}", semester):
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[4], "message": "Untuk status Lulus, isi semester seperti 'Ganjil 2026-2027'."})
            for col in (5, 7, 10):
                if not row[col]:
                    record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[col], "message": "Tanggal wajib diisi untuk status Lulus."})
                elif not _valid_date(raw[col], types[col], workbook.datemode):
                    record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[col], "message": "Gunakan format tanggal dd-mm-yyyy."})
            for col in (6,):
                if not row[col]:
                    record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[col], "message": "Nomor surat yudisium wajib diisi untuk status Lulus."})
            if row[8].casefold() not in _ALLOWED_PATHS:
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[8], "message": "Isi tepat 'Skripsi' atau 'Non Skripsi'."})
            if not _valid_ijazah_number(row[9]):
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[9], "message": "Isi nomor ijazah PDDIKTI/Neo Feeder (bukan kode A/B)."})
        else:
            for col in (5, 7, 10):
                if row[col] and not _valid_date(raw[col], types[col], workbook.datemode):
                    record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[col], "message": "Gunakan format tanggal dd-mm-yyyy."})
            if row[8] and row[8].casefold() not in _ALLOWED_PATHS:
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[8], "message": "Isi tepat 'Skripsi' atau 'Non Skripsi'."})
            if row[9] and not _valid_ijazah_number(row[9]):
                record_errors.append({"row": excel_row, "field": IMPORT_HEADERS[9], "message": "Isi nomor ijazah PDDIKTI/Neo Feeder (bukan kode A/B)."})

        errors.extend(record_errors)
        records.append({"row": excel_row, "nim": nim, "status": row[3], "errors": len(record_errors)})

    workbook.release_resources()
    if not records:
        errors.append({"row": None, "field": "data", "message": "Tidak ada baris data mahasiswa mulai baris 8."})
    return {
        "valid": bool(records) and not errors,
        "total": len(records),
        "valid_rows": sum(1 for record in records if not record["errors"]),
        "invalid_rows": sum(1 for record in records if record["errors"]),
        "preview": records[:20],
        "remaining": max(0, len(records) - 20),
        "errors": errors[:200],
        "warnings": warnings,
    }