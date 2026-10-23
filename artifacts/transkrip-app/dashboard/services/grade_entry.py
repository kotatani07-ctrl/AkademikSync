"""Safe Excel inspection and dry-run validation for eCampuz grade entry."""

from __future__ import annotations

from io import BytesIO
from typing import Any

MAX_ROWS = 10_000
MAX_COLUMNS = 200
MAX_PREVIEW_ROWS = 30
MAX_ERRORS = 200


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
        number_format = cell.number_format
        if number_format and set(number_format) == {"0"}:
            return text.zfill(len(number_format))
        return text
    return str(value).strip()


def _workbook(content: bytes):
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValueError(f"File bukan workbook Excel yang dapat dibaca: {exc}") from exc
    if not workbook.sheetnames:
        workbook.close()
        raise ValueError("Workbook tidak memiliki sheet.")
    return workbook


def inspect_grade_workbook(
    content: bytes,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    """Return sheet/header choices without retaining or modifying uploaded data."""
    if not content:
        raise ValueError("File Excel kosong.")
    if not 1 <= header_row <= MAX_ROWS:
        raise ValueError(f"Baris header harus antara 1 dan {MAX_ROWS}.")

    workbook = _workbook(content)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError("Sheet yang dipilih tidak ada di workbook.")
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ValueError(f"Sheet melebihi batas {MAX_COLUMNS} kolom.")
        rows = sheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            max_col=MAX_COLUMNS,
        )
        header_cells = list(next(rows, ()))
        while header_cells and not _cell_text(header_cells[-1]):
            header_cells.pop()
        columns = [
            {"index": index, "label": _cell_text(cell) or f"Kolom {index + 1}"}
            for index, cell in enumerate(header_cells)
        ]
        if not any(_cell_text(cell) for cell in header_cells):
            raise ValueError(f"Baris header {header_row} kosong.")
        return {
            "sheets": workbook.sheetnames,
            "selected_sheet": sheet.title,
            "header_row": header_row,
            "columns": columns,
        }
    finally:
        workbook.close()


def validate_grade_workbook(
    content: bytes,
    *,
    sheet_name: str,
    header_row: int,
    nim_column: int,
    course_column: int,
    grade_column: int,
) -> dict[str, Any]:
    """Validate mapped rows and produce a dry-run preview; never submits grades."""
    inspected = inspect_grade_workbook(content, sheet_name, header_row)
    column_indexes = {column["index"] for column in inspected["columns"]}
    selected = (nim_column, course_column, grade_column)
    if len(set(selected)) != 3:
        raise ValueError("Kolom NIM, kode mata kuliah, dan nilai harus berbeda.")
    if any(index not in column_indexes for index in selected):
        raise ValueError("Pemetaan kolom tidak sesuai dengan header pada file.")

    workbook = _workbook(content)
    try:
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + MAX_ROWS + 1,
            max_col=max(selected) + 1,
        )
        valid_rows: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        seen_courses: set[tuple[str, str]] = set()
        total_rows = 0
        invalid_rows = 0
        duplicate_rows = 0
        for row_number, cells in enumerate(iterator, start=header_row + 1):
            values = [_cell_text(cells[index]) if index < len(cells) else "" for index in selected]
            nim, course, grade = values
            if not any(values):
                continue
            total_rows += 1
            if total_rows > MAX_ROWS:
                raise ValueError(f"File melebihi batas {MAX_ROWS} baris data.")

            row_errors = []
            if not nim.isdigit() or not 6 <= len(nim) <= 20:
                row_errors.append("NIM harus berupa 6-20 digit angka.")
            if not course:
                row_errors.append("Kode mata kuliah kosong.")
            if not grade:
                row_errors.append("Nilai kosong.")
            if row_errors:
                invalid_rows += 1
                if len(errors) < MAX_ERRORS:
                    errors.append({"row": row_number, "messages": row_errors})
                continue

            key = (nim, course.casefold())
            if key in seen_courses:
                duplicate_rows += 1
                if len(warnings) < MAX_ERRORS:
                    warnings.append({
                        "row": row_number,
                        "message": f"NIM {nim} dan mata kuliah {course} muncul lebih dari sekali; periksa kemungkinan pengulangan.",
                    })
            seen_courses.add(key)
            valid_rows.append({
                "row": row_number,
                "nim": nim,
                "course": course,
                "grade": grade,
            })

        return {
            "total": total_rows,
            "valid": len(valid_rows),
            "invalid": invalid_rows,
            "errors": errors,
            "warnings": warnings,
            "duplicate_rows": duplicate_rows,
            "preview": valid_rows[:MAX_PREVIEW_ROWS],
            "preview_remaining": max(0, len(valid_rows) - MAX_PREVIEW_ROWS),
            "submitted": False,
        }
    finally:
        workbook.close()
