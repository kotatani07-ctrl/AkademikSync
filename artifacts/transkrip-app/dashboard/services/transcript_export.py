"""Build a single-student workbook from the campus transcript workbook."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell


CURRICULA = (
    ("Nilai MHS 2026", "TRANSKRIP NILAI 2026", "2026"),
    ("Nilai MHS 2026 P22", "TRANSKRIP NILAI 22 (2026)", "P22"),
    ("Nilai MHS 2026 P21", "TRANSKRIP NILAI 21 (2026)", "P21"),
)
ROSTER_SHEET = "Mahasiswa 2026"


@dataclass(frozen=True)
class TranscriptExport:
    content: bytes
    filename: str
    sheet_name: str
    curriculum: str


def _cell_nim(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def _clear_row(ws, row: int) -> None:
    for cell in ws[row]:
        if not isinstance(cell, MergedCell):
            cell.value = None


import subprocess
import tempfile


def convert_excel_to_pdf(excel_bytes: bytes, filename_prefix: str = "transkrip") -> bytes:
    """Mengonversi file Excel transkrip ke PDF menggunakan LibreOffice headless."""
    with tempfile.TemporaryDirectory() as tmpdir:
        in_path = Path(tmpdir) / f"{filename_prefix}.xlsx"
        in_path.write_bytes(excel_bytes)
        res = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(in_path)],
            capture_output=True,
            text=True,
        )
        pdf_path = Path(tmpdir) / f"{filename_prefix}.pdf"
        if pdf_path.is_file():
            return pdf_path.read_bytes()
        raise RuntimeError(f"Gagal mengonversi transkrip ke PDF: {res.stderr or res.stdout}")


def build_transcript_export(
    template_path: Path,
    nim: str,
    hasil=None,
    data_pendukung=None,
) -> TranscriptExport:
    """Return a sanitized, formatted workbook for one NIM with support for dynamic injection."""
    nim = _cell_nim(nim)
    if not nim:
        raise ValueError("NIM wajib diisi.")
    if not template_path.is_file():
        raise FileNotFoundError(f"File template Excel tidak ditemukan: {template_path}")

    wb = openpyxl.load_workbook(template_path, data_only=False)
    try:
        # Jika hasil diberikan secara dinamis (misal dari Neo Feeder atau matching engine)
        if hasil is not None and getattr(hasil, 'daftar_mk', None):
            angkatan = getattr(hasil, 'angkatan', None)
            prodi_str = str(getattr(hasil, 'prodi', '')).lower()
            if angkatan == 2021 or 'p21' in prodi_str:
                source_name, transcript_name, curriculum = ("Nilai MHS 2026 P21", "TRANSKRIP NILAI 21 (2026)", "P21")
            elif angkatan == 2022 or 'p22' in prodi_str:
                source_name, transcript_name, curriculum = ("Nilai MHS 2026 P22", "TRANSKRIP NILAI 22 (2026)", "P22")
            else:
                source_name, transcript_name, curriculum = ("Nilai MHS 2026", "TRANSKRIP NILAI 2026", "2026")

            student_name = str(getattr(hasil, 'nama', '') or '').strip() or 'Mahasiswa'
            roster = wb[ROSTER_SHEET]

            # Injeksi identitas mahasiswa di baris 2 (agar selalu masuk ke named range DATA_MHS2026)
            roster.cell(2, 1, 1)  # selector
            roster.cell(2, 3, nim)
            roster.cell(2, 5, student_name)
            roster.cell(2, 6, getattr(hasil, 'tempat_lahir', '') or '')
            roster.cell(2, 7, getattr(hasil, 'tanggal_lahir', '') or '')
            roster.cell(2, 8, (getattr(data_pendukung, 'judul_skripsi', '') or '') if data_pendukung else '')
            roster.cell(2, 9, (getattr(data_pendukung, 'no_ijazah', '') or '') if data_pendukung else '')
            roster.cell(2, 10, getattr(hasil, 'tanggal_lulus', '') or '')

            # Bersihkan baris identitas lainnya
            for row in range(3, roster.max_row + 1):
                _clear_row(roster, row)

            # Injeksi nilai mata kuliah di sheet kurikulum terkait mulai baris 2
            source = wb[source_name]
            course_count = len(hasil.daftar_mk)
            for idx, mk in enumerate(hasil.daftar_mk):
                cr = 2 + idx
                source.cell(cr, 1, idx + 1)
                source.cell(cr, 2, nim)
                source.cell(cr, 3, student_name)
                source.cell(cr, 4, getattr(mk, 'kode_mk', '') or '')
                source.cell(cr, 5, getattr(mk, 'nama_mk', '') or '')
                source.cell(cr, 7, getattr(mk, 'sks', 0) or 0)
                source.cell(cr, 9, f'=IF($J{cr}="","",IF($J{cr}="A",4,IF($J{cr}="A-",3.85,IF($J{cr}="B+",3.25,IF($J{cr}="B",3,IF($J{cr}="B-",2.85,IF($J{cr}="C+",2.35,IF($J{cr}="C",2,0))))))))')
                source.cell(cr, 10, getattr(mk, 'nilai_lambang', '') or '')
                source.cell(cr, 11, f'=IF(I{cr}="","",G{cr}*I{cr})')

            # Bersihkan baris sisa di sheet kurikulum aktif
            for row in range(2 + course_count, source.max_row + 1):
                _clear_row(source, row)

            # Bersihkan sheet kurikulum lainnya dan sheet Kurikulum master
            for other_source, _, _ in CURRICULA:
                if other_source != source_name and other_source in wb.sheetnames:
                    ws_other = wb[other_source]
                    for row in range(2, ws_other.max_row + 1):
                        _clear_row(ws_other, row)

            if "Kurikulum" in wb.sheetnames:
                for row in range(1, wb["Kurikulum"].max_row + 1):
                    _clear_row(wb["Kurikulum"], row)

            # Set selector pada sheet transkrip terkait
            transcript = wb[transcript_name]
            if curriculum == "2026":
                transcript["L2"] = 1
                transcript["M2"] = nim
            elif curriculum == "P21":
                transcript["M2"] = 1
            else:
                transcript["M2"] = f"{nim} - {student_name}"

            # Bersihkan sel selector pada sheet transkrip lainnya
            for _, other_transcript, _ in CURRICULA:
                if other_transcript != transcript_name and other_transcript in wb.sheetnames:
                    for column in range(12, 16):
                        wb[other_transcript].cell(2, column).value = None

            for ws in wb.worksheets:
                ws.sheet_state = "hidden"
            transcript.sheet_state = "visible"
            wb.active = wb.sheetnames.index(transcript_name)

            if getattr(wb, "calculation", None) is not None:
                wb.calculation.fullCalcOnLoad = True
                wb.calculation.forceFullCalc = True
                wb.calculation.calcMode = "auto"

            output = BytesIO()
            wb.save(output)
            safe_curriculum = curriculum.lower().replace(" ", "-")
            return TranscriptExport(
                content=output.getvalue(),
                filename=f"transkrip_{nim}_{safe_curriculum}.xlsx",
                sheet_name=transcript_name,
                curriculum=curriculum,
            )

        # Jalur default jika hasil tidak diberikan (mencari NIM yang sudah ada di workbook)
        matches: list[tuple[str, str, str, list[int]]] = []
        for source_name, transcript_name, curriculum in CURRICULA:
            if source_name not in wb.sheetnames or transcript_name not in wb.sheetnames:
                continue
            source = wb[source_name]
            rows = [
                row
                for row in range(2, source.max_row + 1)
                if _cell_nim(source.cell(row, 2).value) == nim
            ]
            if rows:
                matches.append((source_name, transcript_name, curriculum, rows))

        if not matches:
            raise ValueError(f"NIM {nim} tidak ditemukan pada data kurikulum workbook.")
        if len(matches) > 1:
            raise ValueError(f"NIM {nim} ditemukan pada lebih dari satu kurikulum.")

        source_name, transcript_name, curriculum, course_rows = matches[0]
        if ROSTER_SHEET not in wb.sheetnames:
            raise ValueError("Sheet identitas mahasiswa tidak ditemukan pada workbook.")
        roster = wb[ROSTER_SHEET]
        roster_row = next(
            (
                row
                for row in range(2, roster.max_row + 1)
                if _cell_nim(roster.cell(row, 3).value) == nim
            ),
            None,
        )
        if roster_row is None:
            raise ValueError(f"NIM {nim} belum memiliki data identitas di workbook.")

        selector = roster.cell(roster_row, 1).value
        student_name = str(roster.cell(roster_row, 5).value or "").strip()
        if selector is None or not student_name:
            raise ValueError(f"Identitas NIM {nim} belum lengkap di workbook.")

        for ws in wb.worksheets:
            if ws.title == "Kurikulum":
                for row in range(1, ws.max_row + 1):
                    _clear_row(ws, row)
            elif ws.title == ROSTER_SHEET:
                for row in range(2, ws.max_row + 1):
                    if row != roster_row:
                        _clear_row(ws, row)
            elif ws.title in {item[0] for item in CURRICULA}:
                for row in range(2, ws.max_row + 1):
                    keep = ws.title == source_name and row in course_rows
                    if not keep:
                        _clear_row(ws, row)

        transcript = wb[transcript_name]
        if curriculum == "2026":
            transcript["L2"] = selector
        elif curriculum == "P21":
            transcript["M2"] = selector
        else:
            transcript["M2"] = f"{nim} - {student_name}"

        for _, other_transcript, _ in CURRICULA:
            if other_transcript != transcript_name and other_transcript in wb.sheetnames:
                for column in range(12, 16):
                    wb[other_transcript].cell(2, column).value = None

        for ws in wb.worksheets:
            ws.sheet_state = "hidden"
        transcript.sheet_state = "visible"
        wb.active = wb.sheetnames.index(transcript_name)
        if getattr(wb, "calculation", None) is not None:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"

        output = BytesIO()
        wb.save(output)
        safe_curriculum = curriculum.lower().replace(" ", "-")
        return TranscriptExport(
            content=output.getvalue(),
            filename=f"transkrip_{nim}_{safe_curriculum}.xlsx",
            sheet_name=transcript_name,
            curriculum=curriculum,
        )
    finally:
        wb.close()