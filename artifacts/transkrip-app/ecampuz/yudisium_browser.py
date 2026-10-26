"""Guarded, sequential browser workflow for yudisium registration."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urljoin, urlsplit

from .config import Config, STATE_DIR
from .credentials import CredentialStore, write_hidden



def _cell_text(value: Any, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and re.fullmatch(r"0{2,}", number_format or ""):
        return f"{value:0{len(number_format)}d}"
    return str(value).strip()


def _normalize_headers(headers: list[Any]) -> list[str]:
    normalized = [str(value or "").strip() for value in headers]
    if not normalized or any(not value for value in normalized):
        raise ValueError("Semua kolom input harus memiliki nama header.")
    lowered = [value.casefold() for value in normalized]
    if len(lowered) != len(set(lowered)):
        raise ValueError("Nama header input tidak boleh duplikat.")
    return normalized


def read_records(path: Path, nim_column: str, sheet: str | None = None) -> list[dict[str, str]]:
    """Read and validate student rows from CSV or XLSX."""
    if not path.is_file():
        raise FileNotFoundError(f"File input tidak ditemukan: {path}")

    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(4096)
            stream.seek(0)
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = ","
            reader = csv.DictReader(stream, delimiter=delimiter)
            original_headers = reader.fieldnames or []
            headers = _normalize_headers(original_headers)
            source_rows = [
                {
                    header: _cell_text(row.get(original))
                    for header, original in zip(headers, original_headers)
                }
                for row in reader
            ]
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        import openpyxl

        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            if sheet:
                if sheet not in workbook.sheetnames:
                    raise ValueError(f"Sheet {sheet!r} tidak ditemukan: {workbook.sheetnames}")
                worksheet = workbook[sheet]
            else:
                worksheet = workbook.active
            rows = worksheet.iter_rows()
            header_cells = next(rows, None)
            headers = _normalize_headers(
                [cell.value for cell in header_cells] if header_cells else []
            )
            source_rows = []
            for cells in rows:
                source_rows.append(
                    {
                        header: _cell_text(cells[index].value, cells[index].number_format)
                        if index < len(cells)
                        else ""
                        for index, header in enumerate(headers)
                    }
                )
        finally:
            workbook.close()
    else:
        raise ValueError("Format input harus .csv, .xlsx, atau .xlsm.")

    header_lookup = {header.casefold(): header for header in headers}
    actual_nim_column = header_lookup.get(nim_column.casefold())
    if not actual_nim_column:
        raise ValueError(
            f"Kolom NIM {nim_column!r} tidak ditemukan. Header tersedia: {headers}"
        )

    records: list[dict[str, str]] = []
    seen: set[str] = set()
    for row_number, row in enumerate(source_rows, start=2):
        if not any(value.strip() for value in row.values()):
            continue
        nim = row.get(actual_nim_column, "").strip()
        if not nim:
            raise ValueError(f"NIM kosong pada baris {row_number}.")
        key = nim.casefold()
        if key in seen:
            raise ValueError(f"NIM duplikat {nim!r} pada baris {row_number}.")
        seen.add(key)
        row[actual_nim_column] = nim
        records.append(row)

    if not records:
        raise ValueError("Tidak ada baris mahasiswa untuk diproses.")
    return records


def inspect_forms(page) -> list[dict[str, Any]]:
    """Return form structure without exposing hidden token values."""
    return page.locator("form").evaluate_all(
        """forms => forms.map((form, index) => ({
          index,
          action: form.getAttribute('action') || '',
          method: (form.getAttribute('method') || 'get').toLowerCase(),
          controls: Array.from(form.querySelectorAll('input, select, textarea, button'))
            .map(el => ({
              tag: el.tagName.toLowerCase(),
              name: el.name || '',
              id: el.id || '',
              type: (el.type || '').toLowerCase(),
              required: Boolean(el.required),
              disabled: Boolean(el.disabled),
              placeholder: el.getAttribute('placeholder') || '',
              label: el.labels && el.labels.length
                ? Array.from(el.labels).map(label => label.innerText.trim()).join(' ')
                : (el.getAttribute('aria-label') || ''),
              options: el.tagName === 'SELECT'
                ? Array.from(el.options).map(option => ({
                    value: option.value,
                    label: option.text.trim()
                  }))
                : []
            }))
        }))"""
    )


def _is_login_page(page) -> bool:
    return page.locator("#aid_username, #gsFormLogin, #gsformlogin, input[name='user[passwd]']").count() > 0


def _manual_login(
    page,
    base_url: str,
    username: str,
    password: str,
    pause_callback=None,
) -> str | None:
    logged_in = False
    if not page.url.startswith(base_url):
        page.goto(base_url, wait_until="domcontentloaded")
    if _is_login_page(page):
        logged_in = True
        if username and password:
            try:
                from .browser_login import _fill_login_credentials
                _fill_login_credentials(page, username, password)
            except Exception:
                pass
        print("Login dan CAPTCHA secara manual di jendela browser.")
        print("Sesudah login, buka halaman Tambah - Mahasiswa Lain untuk sesi aktif.")
        if pause_callback:
            pause_callback(page)
        else:
            input("Setelah halaman tersebut terbuka, tekan Enter di terminal ini: ")
        try:
            page.wait_for_load_state("domcontentloaded", timeout=10_000)
        except Exception:
            pass
    if _is_login_page(page):
        raise RuntimeError("Halaman login masih tampil; sesi browser belum aktif.")
    current = urlsplit(page.url)
    base = urlsplit(base_url)
    if current.scheme != base.scheme or current.netloc.casefold() != base.netloc.casefold():
        raise RuntimeError("Browser harus tetap berada pada origin eCampuz setelah login.")
    if logged_in and "/kelulusan/" not in current.path.casefold():
        raise RuntimeError("Setelah login, buka halaman Tambah - Mahasiswa Lain sebelum menekan Enter.")
    return page.url if logged_in else None


def _find_nim_range(page):
    rows = page.locator("tr")
    for index in range(rows.count()):
        row = rows.nth(index)
        text = re.sub(r"\s+", " ", row.inner_text()).strip()
        if not re.search(r"\bNIM\b", text, re.I):
            continue
        inputs = row.locator("input[type='text'], input:not([type])")
        visible = [
            inputs.nth(input_index)
            for input_index in range(inputs.count())
            if inputs.nth(input_index).is_visible() and inputs.nth(input_index).is_enabled()
        ]
        if len(visible) >= 2:
            return visible[0], visible[1]
    raise ValueError(
        "Tidak menemukan dua input rentang NIM pada baris filter NIM. "
        "Pastikan URL berada di halaman Tambah/Pendaftaran Yudisium, bukan halaman menu."
    )


def _period_key(value: str) -> str:
    """Normalisasi label periode eCampuz, mis. Genap 2022-2023 / 20222."""
    text = re.sub(r"\s+", " ", str(value or "")).strip().casefold()
    text = text.replace("gasal", "ganjil")
    text = text.replace("–", "-").replace("—", "-").replace("-", "/")
    match = re.search(r"\b(ganjil|genap|pendek)\b\D*(\d{4})\D+(\d{4})", text)
    if match:
        return f"{match.group(1)} {match.group(2)}/{match.group(3)}"
    compact = re.sub(r"\s+", "", text)
    match = re.fullmatch(r"(20\d{2})([12])", compact)
    if match:
        year = int(match.group(1))
        term = "ganjil" if match.group(2) == "1" else "genap"
        return f"{term} {year}/{year + 1}"
    return text


def _select_semester(page, period: str) -> str | None:
    """Pilih dropdown Semester eCampuz sesuai ID/label periode."""
    if not period:
        return None

    selects = page.locator("select[name='semester']")
    target_select = None
    for index in range(selects.count()):
        candidate = selects.nth(index)
        if candidate.is_visible():
            target_select = candidate
            break
    if target_select is None:
        raise RuntimeError(
            f"Dropdown Semester untuk periode {period!r} tidak ditemukan; "
            "pendaftaran tidak dilanjutkan agar tidak masuk semester yang salah."
        )

    choices: list[tuple[str, str]] = []
    options = target_select.locator("option")
    for index in range(options.count()):
        option = options.nth(index)
        label = re.sub(r"\s+", " ", option.inner_text()).strip()
        if label:
            choices.append((label, _period_key(label)))

    target_key = _period_key(period)
    chosen = next((label for label, key in choices if key == target_key), None)
    if chosen is None:
        # Fallback untuk label yang sudah sama persis setelah normalisasi.
        normalized = re.sub(r"\s+", " ", period).strip().casefold().replace("-", "/")
        chosen = next(
            (
                label
                for label, _ in choices
                if re.sub(r"\s+", " ", label).strip().casefold().replace("-", "/")
                == normalized
            ),
            None,
        )
    if chosen is None:
        available = ", ".join(label for label, _ in choices)
        raise ValueError(
            f"Semester {period!r} tidak ditemukan pada halaman eCampuz. "
            f"Pilihan yang tersedia: {available}"
        )

    target_select.select_option(label=chosen)
    return chosen


def _is_registration_page(page) -> bool:
    """True bila halaman aktif adalah form Tambah/Pendaftaran Yudisium."""
    body = page.locator("body").inner_text()
    is_add_page = (
        "Tambah Mahasiswa Lain" in body
        or "Tambah - Pilih Mahasiswa" in body
        or "Tambah - Mahasiswa Lain" in body
    )
    has_nim_filter = (
        page.locator("input[name='filter_awal']").count() > 0
        and page.locator("input[name='filter_akhir']").count() > 0
    )
    return is_add_page and has_nim_filter


def _discover_registration_page(page, base_url: str, timeout_ms: int = 60000) -> str | None:
    """Cari ulang URL Tambah Yudisium pada sesi eCampuz yang sedang aktif.

    URL menu eCampuz memakai token ``act`` yang dapat berubah setelah login
    ulang. Jangan mengandalkan URL yang tersimpan di web.json; bila URL itu
    sudah kedaluwarsa, halaman hanya menampilkan menu. Navigasi lewat teks menu
    membuat alur tetap aman tanpa menebak token.
    """
    if _is_registration_page(page):
        return page.url

    base_parts = urlsplit(base_url)
    current = urlsplit(page.url)
    if (
        current.scheme.casefold() != base_parts.scheme.casefold()
        or current.netloc.casefold() != base_parts.netloc.casefold()
    ):
        page.goto(base_url, wait_until="domcontentloaded", timeout=timeout_ms)
        if _is_login_page(page):
            return None

    menu_url = _link_url_by_text(page, "Pendaftaran Yudisium")
    if menu_url:
        try:
            menu_url = _ensure_same_origin(menu_url, base_url)
        except ValueError:
            menu_url = None
    if menu_url:
        page.goto(menu_url, wait_until="domcontentloaded", timeout=timeout_ms)
        try:
            page.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
        except Exception:
            pass
        if _is_login_page(page):
            return None

    if _is_registration_page(page):
        return page.url

    # Halaman daftar Pendaftaran Yudisium memiliki tautan "Tambah"; halaman
    # Tambah - Pilih Mahasiswa memiliki tautan "Tambah Mahasiswa Lain".
    for label in (
        "Tambah",
        "Tambah Mahasiswa Lain",
        "Tambah - Pilih Mahasiswa",
        "Tambah - Mahasiswa Lain",
    ):
        link = _link_url_by_text(page, label)
        if not link:
            continue
        try:
            link = _ensure_same_origin(link, base_url)
        except ValueError:
            continue
        page.goto(link, wait_until="domcontentloaded", timeout=timeout_ms)
        try:
            page.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
        except Exception:
            pass
        if _is_login_page(page):
            return None
        if _is_registration_page(page):
            return page.url

    return None


def _ensure_registration_page(page) -> None:
    """Pastikan URL adalah halaman tambah, bukan menu atau daftar pendaftaran."""
    if not _is_registration_page(page):
        raise RuntimeError(
            "Halaman yang dibuka bukan halaman Tambah/Pendaftaran Yudisium. "
            "Tempel URL halaman Tambah - Mahasiswa Lain yang aktif pada sesi "
            "eCampuz ini, atau URL menu Pendaftaran Yudisium agar URL aktual "
            "ditemukan otomatis."
        )


def _registration_detail_submit(page, nim: str):
    """Kembalikan tombol Tambah jika halaman detail pendaftaran sudah terbuka."""
    body = page.locator("body").inner_text()
    if "Pendaftaran Yudisium Mahasiswa" not in body:
        return None

    niu = page.locator("input[name='mhsNiu']")
    if niu.count() == 0:
        return None
    detail_nim = niu.nth(0).input_value().strip()
    if detail_nim and detail_nim != nim:
        raise RuntimeError(
            f"Halaman detail menampilkan NIM {detail_nim}, bukan NIM yang diminta {nim}."
        )

    yds_id = page.locator("input[name='ydsId']")
    if yds_id.count() and yds_id.nth(0).input_value().strip():
        raise RuntimeError(
            f"NIM {nim} sudah memiliki pendaftaran yudisium "
            f"(ydsId={yds_id.nth(0).input_value().strip()}); tidak diulang."
        )

    try:
        return _button_by_text(page, "Tambah")
    except ValueError:
        return None


def _parse_registration_date(value: str) -> tuple[str, str, str] | None:
    """Validasi tanggal daftar DD-MM-YYYY dan kembalikan bagian tanggal."""
    text = str(value or "").strip()
    match = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", text)
    if not match:
        return None
    day, month, year = match.groups()
    try:
        date(int(year), int(month), int(day))
    except ValueError:
        return None
    return f"{int(day):02d}", f"{int(month):02d}", year


def _fill_registration_date(page, value: str) -> None:
    """Isi Tanggal Daftar eCampuz; field ini wajib agar Tambah diproses server."""
    text = str(value or "").strip()
    if not text:
        return
    parts = _parse_registration_date(text)
    if parts is None:
        raise ValueError(
            f"Tanggal daftar {text!r} tidak valid; gunakan format DD-MM-YYYY, "
            "misalnya 31-03-2023."
        )

    day, month, year = parts
    expected = f"{day}-{month}-{year}"
    fields = page.locator("input[name='tgldaftar']")
    if fields.count() == 0:
        raise RuntimeError("Field Tanggal Daftar tidak ditemukan pada halaman detail.")
    target = fields.nth(0)
    page.evaluate(
        """([expected, day, month, year]) => {
          const setValue = (name, value) => {
            const el = document.querySelector(`input[name='${name}']`);
            if (el) el.value = value;
          };
          setValue('tgldaftar_day', day);
          setValue('tgldaftar_mon', month);
          setValue('tgldaftar_year', year);
          const el = document.querySelector("input[name='tgldaftar']");
          if (el) {
            el.value = expected;
            el.dispatchEvent(new Event('input', {bubbles: true}));
            el.dispatchEvent(new Event('change', {bubbles: true}));
          }
          // Jangan biarkan datepicker menutupi tombol Tambah.
          const picker = document.getElementById('ui-datepicker-div');
          if (picker) picker.style.display = 'none';
        }""",
        [expected, day, month, year],
    )
    actual = target.input_value().strip()
    if actual != expected:
        raise RuntimeError(
            f"Gagal mengisi Tanggal Daftar: nilai terbaca {actual!r}, "
            f"seharusnya {expected!r}."
        )


def _save_uncertain_artifact(page, nim: str) -> Path | None:
    """Simpan screenshot/HTML saat hasil tidak pasti untuk pemeriksaan manual."""
    try:
        out_dir = STATE_DIR / "yudisium_uncertain"
        out_dir.mkdir(parents=True, exist_ok=True)
        base = out_dir / f"{nim}_{time.strftime('%Y%m%d-%H%M%S')}"
        page.screenshot(path=str(base.with_suffix(".png")), full_page=True)
        base.with_suffix(".html").write_text(page.content(), encoding="utf-8")
        return base
    except Exception:
        return None


def _link_url_by_text(page, label: str) -> str | None:
    links = page.locator("a[href]")
    for index in range(links.count()):
        link = links.nth(index)
        text = re.sub(r"\s+", " ", link.inner_text()).strip()
        if text.casefold() != label.casefold():
            continue
        href = (link.get_attribute("href") or "").strip()
        if href and not href.startswith(("#", "javascript:")):
            return urljoin(page.url, href)
    return None


def _management_has_registration(
    page,
    nim: str,
    period: str,
    timeout_ms: int,
) -> bool | None:
    """Cek NIM pada menu Pengelolaan Yudisium untuk semester yang dipilih.

    Halaman Daftar Pendaftaran Yudisium mengikuti semester global dan bisa
    menampilkan "Data tidak ditemukan" meskipun pendaftaran berhasil. Menu
    Pengelolaan Yudisium memiliki filter semester, jadi dipakai untuk
    verifikasi sebelum/sesudah submit.
    """
    href = _link_url_by_text(page, "Pengelolaan Yudisium")
    if not href:
        return None
    try:
        href = _ensure_same_origin(href, page.url)
    except ValueError:
        return None

    verifier = page.context.new_page()
    try:
        verifier.goto(href, wait_until="domcontentloaded", timeout=timeout_ms)
        try:
            verifier.wait_for_load_state(
                "networkidle", timeout=min(timeout_ms, 15000)
            )
        except Exception:
            pass
        if _is_login_page(verifier):
            return None

        _select_semester(verifier, period)
        verifier.locator("input[name='filter_awal']").fill(nim)
        verifier.locator("input[name='filter_akhir']").fill(nim)
        _button_by_text(verifier, "TAMPILKAN").click()
        try:
            verifier.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
        except Exception:
            pass
        # Jangan memutuskan "tidak terdaftar" hanya dari jeda 500 ms; tabel
        # eCampuz dapat merender hasil setelah request AJAX selesai.
        try:
            _wait_for_results(verifier, nim, min(timeout_ms, 10000))
        except Exception:
            pass
        return _has_exact_student_row(verifier, nim)
    finally:
        try:
            verifier.close()
        except Exception:
            pass


def _button_by_text(page, label: str):
    candidates = page.locator(
        "button, input[type='submit'], input[type='button'], input[type='reset']"
    )
    matches = []
    for index in range(candidates.count()):
        candidate = candidates.nth(index)
        text = candidate.inner_text() if candidate.evaluate("el => el.tagName.toLowerCase() === 'button'") else ""
        text = text or candidate.get_attribute("value") or ""
        if re.sub(r"\s+", " ", text).strip().casefold() == label.casefold() and candidate.is_visible():
            matches.append(candidate)
    if len(matches) >= 1:
        return matches[0]
    # Fallback pencarian substring jika exact match tidak ditemukan
    for index in range(candidates.count()):
        candidate = candidates.nth(index)
        text = candidate.inner_text() if candidate.evaluate("el => el.tagName.toLowerCase() === 'button'") else ""
        text = text or candidate.get_attribute("value") or candidate.get_attribute("title") or ""
        if label.casefold() in re.sub(r"\s+", " ", text).strip().casefold() and candidate.is_visible():
            matches.append(candidate)
    if len(matches) >= 1:
        return matches[0]
    raise ValueError(f"Tombol {label!r} tidak ditemukan pada halaman.")


def _try_reset_filter(page) -> bool:
    """Mencoba mereset filter pencarian jika tombol reset tersedia.
    Jika tidak ada tombol reset atau gagal di-klik, fungsi ini tidak melempar error."""
    try:
        candidates = page.locator(
            "button, input[type='submit'], input[type='button'], input[type='reset'], a.btn"
        )
        for index in range(candidates.count()):
            candidate = candidates.nth(index)
            if not candidate.is_visible():
                continue
            text = candidate.inner_text() if candidate.evaluate("el => el.tagName.toLowerCase() in ['button', 'a']") else ""
            text = text or candidate.get_attribute("value") or candidate.get_attribute("title") or ""
            norm_text = re.sub(r"\s+", " ", text).strip().casefold()
            if norm_text in ("reset", "bersihkan", "batal", "reset filter"):
                candidate.click()
                try:
                    page.wait_for_load_state("domcontentloaded", timeout=2000)
                except Exception:
                    pass
                return True
    except Exception:
        pass
    return False


def _wait_for_results(page, nim: str, timeout_ms: int) -> None:
    page.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
    no_data_pattern = r"data\s+tidak\s+ditemukan|tidak\s+ada\s+data|no\s+data\s+found"
    try:
        # Tunggu NIM ATAU pesan "data tidak ditemukan" agar NIM yang memang
        # tidak ada tidak membuat bot menunggu timeout penuh.
        page.wait_for_function(
            """(args) => {
              const text = document.body ? document.body.innerText : '';
              return text.includes(args.nim) || new RegExp(args.pattern, 'i').test(text);
            }""",
            arg={"nim": nim, "pattern": no_data_pattern},
            timeout=timeout_ms,
        )
    except Exception:
        raise TimeoutError(f"Hasil pencarian NIM {nim} tidak muncul sebelum timeout.")
    try:
        page.get_by_text(nim, exact=True).first.wait_for(state="visible", timeout=2000)
        return
    except Exception:
        pass
    body = page.locator("body").inner_text()
    if re.search(no_data_pattern, body, re.I):
        raise LookupError(f"NIM {nim} tidak ditemukan pada hasil pencarian (Data tidak ditemukan).")
    raise TimeoutError(f"Hasil pencarian NIM {nim} tidak muncul sebelum timeout.")


def _has_exact_student_row(page, nim: str) -> bool:
    rows = page.locator("table tr")
    for index in range(rows.count()):
        cells = rows.nth(index).locator("th, td")
        for cell_index in range(cells.count()):
            if re.sub(r"\s+", " ", cells.nth(cell_index).inner_text()).strip() == nim:
                return True
    return False


def _find_exact_student_row(page, nim: str):
    rows = page.locator("table tr")
    matches = []
    for index in range(rows.count()):
        row = rows.nth(index)
        cells = row.locator("th, td")
        for cell_index in range(cells.count()):
            if re.sub(r"\s+", " ", cells.nth(cell_index).inner_text()).strip() == nim:
                matches.append(row)
                break
    if len(matches) > 1:
        action_rows = [
            r for r in matches
            if r.locator("a:visible, button:visible, input[type='button']:visible, input[type='submit']:visible").count() > 0
        ]
        if len(action_rows) == 1:
            return action_rows[0]
    if len(matches) != 1:
        raise LookupError(f"NIM {nim} menghasilkan {len(matches)} baris yang cocok persis.")
    return matches[0]


def _student_action(row):
    actions = row.locator("a:visible, button:visible, input[type='button']:visible, input[type='submit']:visible")
    if actions.count() < 1:
        raise LookupError("Baris mahasiswa tidak memiliki tombol/link aksi yang terlihat.")

    if actions.count() == 1:
        return actions.first

    scored_candidates = []
    for index in range(actions.count()):
        action = actions.nth(index)
        inner_text = action.inner_text().strip().casefold()
        val = (action.get_attribute("value") or "").strip().casefold()
        title = (action.get_attribute("title") or "").strip().casefold()
        aria_label = (action.get_attribute("aria-label") or "").strip().casefold()
        cls = (action.get_attribute("class") or "").strip().casefold()
        onclick = (action.get_attribute("onclick") or "").strip().casefold()
        href = (action.get_attribute("href") or "").strip().casefold()

        # Ambil class dari elemen ikon turunan (i, svg, span)
        icon_classes = ""
        try:
            icon_classes = action.evaluate(
                """el => Array.from(el.querySelectorAll('i, svg, span'))
                              .map(i => i.className || '')
                              .join(' ')"""
            ).casefold()
        except Exception:
            pass

        metadata_text = f"{inner_text} {val} {title} {aria_label} {cls} {onclick} {href} {icon_classes}"

        score = 0
        # Sinyal positif untuk aksi pendaftaran / pemilihan mahasiswa yudisium
        if re.search(r"\b(daftar|daftarkan|registrasi|register|pilih|tambah|tambahkan|pilihkan|enroll|select)\b", metadata_text):
            score += 20
        if re.search(r"fa-plus|fa-user-plus|fa-check|fa-user-check|glyphicon-plus|glyphicon-ok|bi-plus|bi-check", metadata_text):
            score += 15
        if re.search(r"btn-success|btn-primary", cls):
            score += 5
        if re.search(r"(?:daftar|pilih|tambah|register|add)", href + onclick):
            score += 10

        # Sinyal negatif untuk aksi melihat detail / profil / riwayat / hapus
        if re.search(r"\b(detail|lihat|view|rincian|profil|profile|riwayat|history|biodata|foto|photo|info|transkrip|krs)\b", metadata_text):
            score -= 25
        if re.search(r"fa-eye|fa-info|fa-magnifying-glass|glyphicon-eye|btn-info", metadata_text):
            score -= 20
        if re.search(r"\b(hapus|delete|batal|cancel|remove)\b|fa-trash|fa-times|fa-remove|btn-danger", metadata_text):
            score -= 30

        scored_candidates.append((score, index, action, metadata_text))

    # Urutkan kandidat berdasarkan skor tertinggi
    scored_candidates.sort(key=lambda item: item[0], reverse=True)
    best_score, best_index, best_action, best_meta = scored_candidates[0]

    # 1. Jika kandidat terbaik memiliki skor positif yang jelas
    if best_score > 0:
        runner_up = scored_candidates[1] if len(scored_candidates) > 1 else None
        if runner_up and runner_up[0] == best_score:
            # Jika skor seri positif, prioritaskan elemen button/input
            for s, idx, act, meta in scored_candidates:
                if s == best_score and act.evaluate("el => el.tagName.toLowerCase() in ['button', 'input']"):
                    return act
        return best_action

    # 2. Jika tidak ada skor positif, namun ada kandidat non-negatif
    non_negative = [item for item in scored_candidates if item[0] >= 0]
    if len(non_negative) == 1:
        return non_negative[0][2]

    # 3. Fallback: cari tombol dengan ikon search / lookup jika ada tepat 1
    try:
        icon_indexes = actions.evaluate_all(
            """nodes => nodes.flatMap((node, index) => {
              const icon = node.querySelector(
                            '[class~="fa-search"], [class~="fa-magnifying-glass"], ' +
                            '[class~="glyphicon-search"], [class~="bi-search"], ' +
                            '[data-icon="search"], [aria-label*="search" i]'
              );
              if (!icon) return [];
              const owner = icon.closest('a, button, input[type="button"], input[type="submit"]');
              return owner === node ? [index] : [];
            })"""
        )
        if len(icon_indexes) == 1:
            return actions.nth(icon_indexes[0])
    except Exception:
        pass

    # 4. Jika kandidat terbaik lebih baik daripada yang lain
    if len(scored_candidates) > 1 and best_score > scored_candidates[1][0]:
        return best_action

    raise LookupError(
        f"Baris mahasiswa memiliki {actions.count()} aksi dan tidak dapat ditentukan dengan aman. "
        f"Kandidat: {[meta[:60] for _, _, _, meta in scored_candidates]}"
    )


def _classify_registration_result(
    dialog_messages: list[str],
    notice_messages: list[str],
    body_text: str,
) -> tuple[str, str]:
    rejection = re.compile(
        r"kurang|gagal|error|ditolak|tidak dapat|tidak bisa|belum memenuhi|tidak memenuhi",
        re.I,
    )
    success = re.compile(
        r"berhasil\s+(?:didaftarkan|ditambahkan|disimpan)|pendaftaran\s+berhasil|"
        r"sukses\s+(?:didaftarkan|ditambahkan)|telah\s+terdaftar|successfully\s+registered",
        re.I,
    )
    messages = [message.strip() for message in [*dialog_messages, *notice_messages] if message.strip()]
    for message in messages:
        if rejection.search(message):
            return "failed", message[:500]
    for message in messages:
        if success.search(message) and not re.search(r"apakah|yakin", message, re.I):
            return "submitted", message[:500]
    if rejection.search(body_text):
        return "failed", next(
            (line.strip()[:500] for line in body_text.splitlines() if rejection.search(line)),
            "Pendaftaran ditolak oleh server.",
        )
    if success.search(body_text):
        return "submitted", next(
            (line.strip()[:500] for line in body_text.splitlines() if success.search(line)),
            "Pendaftaran berhasil dikonfirmasi oleh halaman.",
        )
    return "uncertain", "Tidak ditemukan konfirmasi sukses atau penolakan yang jelas."


def _registration_notices(page) -> list[str]:
    return page.locator(".alert, .toast, [role='alert'], .swal2-html-container").evaluate_all(
        "nodes => nodes.map(node => (node.innerText || '').trim()).filter(Boolean)"
    )


def _log_result(path: Path, nim: str, status: str, reason: str = "") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    write_header = not path.exists() or path.stat().st_size == 0
    with path.open("a", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        if write_header:
            writer.writerow(["waktu", "nim", "status", "keterangan"])
        writer.writerow([time.strftime("%Y-%m-%dT%H:%M:%S"), nim, status, reason])


def _resolve_profile_credentials(
    config: Config,
    *,
    allow_prompt: bool = True,
) -> tuple[str, str]:
    """Resolve eAkademik credentials without blocking non-interactive (web) jobs.

    The web worker runs inside the Django process, whose stdin may still be a TTY
    when runserver is started from a terminal. Prompting there blocks the job
    forever before Chromium is launched. In non-interactive mode we deliberately
    skip the prompt and let the browser continue with the saved session cookies.
    """
    try:
        store = CredentialStore(config.credential_path)
        credentials = store.load()
        if credentials:
            return credentials
    except Exception:
        pass

    if not allow_prompt:
        print(
            "[info] Mode web/non-interaktif: tidak meminta kredensial eAkademik "
            "di terminal; memakai sesi/cookies yang tersimpan."
        )
        return ("", "")

    import sys
    if getattr(sys, "stdin", None) and hasattr(sys.stdin, "isatty") and sys.stdin.isatty():
        try:
            import getpass

            username = input("Username eAkademik: ").strip()
            password = getpass.getpass("Password eAkademik: ")
            if username and password:
                try:
                    store.save(username, password)
                except Exception:
                    pass
                return username, password
        except (OSError, EOFError):
            pass
    return ("", "")


def _inject_session_cookies(context, config: Config) -> bool:
    session_file = config.session_path
    if not session_file.exists():
        return False
    try:
        session_data = json.loads(session_file.read_text(encoding="utf-8"))
        cookies_dict = session_data.get("cookies", {})
        if not cookies_dict:
            return False
        base_url = session_data.get("base_url") or config.base_url
        playwright_cookies = [
            {"name": str(name), "value": str(value), "url": base_url}
            for name, value in cookies_dict.items()
        ]
        if playwright_cookies:
            context.add_cookies(playwright_cookies)
            return True
    except Exception:
        pass
    return False


def _ensure_same_origin(target_url: str, base_url: str) -> str:
    target = urlsplit(target_url)
    base = urlsplit(base_url)
    if target.scheme != base.scheme or target.netloc.casefold() != base.netloc.casefold():
        raise ValueError("URL pendaftaran harus berada pada origin eCampuz yang dikonfigurasi.")
    return target_url


def _campaign_fingerprint(input_path: Path, base_url: str, campaign_id: str) -> str:
    digest = hashlib.sha256()
    digest.update(base_url.encode("utf-8"))
    digest.update(campaign_id.encode("utf-8"))
    digest.update(input_path.read_bytes())
    return digest.hexdigest()


def _load_state(path: Path, fingerprint: str) -> dict[str, Any]:
    if not path.exists():
        return {"fingerprint": fingerprint, "rows": {}}
    try:
        state = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Checkpoint yudisium tidak dapat dibaca: {exc}") from exc
    if state.get("fingerprint") != fingerprint:
        raise ValueError(
            f"Checkpoint {path} berasal dari file atau kampus berbeda. "
            "Gunakan --state-file baru setelah memeriksa pendaftaran yang sudah terkirim."
        )
    return state


def _pending_records(
    records: list[dict[str, str]],
    nim_column: str,
    state: dict[str, Any],
    retry_failed: bool,
) -> list[dict[str, str]]:
    statuses = state.get("rows", {})
    uncertain = [
        nim for nim, item in statuses.items()
        if item.get("status") in {"submitting", "uncertain"}
    ]
    if uncertain:
        raise RuntimeError(
            "Checkpoint memiliki hasil belum pasti untuk NIM: "
            + ", ".join(uncertain[:10])
            + ". Verifikasi di eCampuz secara manual; bot tidak akan mengulangnya."
        )
    skipped_statuses = {"submitted"} if retry_failed else {"submitted", "failed"}
    return [
        record for record in records
        if statuses.get(record[nim_column], {}).get("status") not in skipped_statuses
    ]


def _save_state(path: Path, state: dict[str, Any]) -> None:
    write_hidden(path, json.dumps(state, indent=2, ensure_ascii=False) + "\n")


def _emit(args, message: str) -> None:
    callback = getattr(args, "on_log", None)
    if callback:
        callback(message)
    else:
        print(message)


def _print_inspection(page) -> None:
    forms = inspect_forms(page)
    tables = page.locator("table").evaluate_all(
        """tables => tables.map((table, index) => {
          const rows = Array.from(table.querySelectorAll('tr'));
          return {
            index,
            rows: rows.length,
            headers: rows.length
              ? Array.from(rows[0].querySelectorAll('th, td')).map(cell => cell.innerText.trim())
              : []
          };
        })"""
    )
    links = page.locator("a[href]").evaluate_all(
        """nodes => nodes.map(node => ({
          text: (node.innerText || '').trim().replace(/\\s+/g, ' '),
          href: node.href
        })).filter(link => link.text || link.href)"""
    )
    buttons = page.locator("button, input[type=submit], input[type=button]").evaluate_all(
        """nodes => nodes.map(node => ({
          tag: node.tagName.toLowerCase(),
          type: (node.type || '').toLowerCase(),
          name: node.name || '',
          id: node.id || '',
          text: (node.innerText || node.value || '').trim(),
          disabled: Boolean(node.disabled)
        }))"""
    )
    result = {
        "url": page.url,
        "title": page.title(),
        "forms": forms,
        "tables": tables,
        "links": links,
        "buttons": buttons,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if not forms:
        print("[info] Tidak ada form HTML pada halaman ini.")


def run_yudisium(args, config: Config) -> int:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError(
            "Playwright belum terpasang. Jalankan pip install -r requirements.txt "
            "lalu python -m playwright install chromium."
        ) from exc

    records: list[dict[str, str]] = []
    input_path: Path | None = None
    nim_column = args.nim_column
    if not args.inspect:
        if not args.file:
            raise ValueError("Mode batch memerlukan --file berisi NIM; gunakan --inspect untuk inspeksi form.")
        input_path = Path(args.file).resolve()
        records = read_records(input_path, args.nim_column, sheet=args.sheet)
        nim_column = next(
            name for name in records[0]
            if name.casefold() == args.nim_column.casefold()
        )
        if args.limit is not None:
            if args.limit < 1:
                raise ValueError("--limit harus lebih besar dari nol.")
        if args.submit and args.limit is None:
            raise ValueError("Submit wajib memakai --limit sebagai batas eksplisit.")
        if not args.submit:
            preview = records[: args.limit] if args.limit else records
            _emit(args, f"[dry-run] {len(records)} baris valid; tidak ada browser yang dibuka dan tidak ada submit.")
            _emit(args, f"[dry-run] Kolom: {list(records[0])}")
            _emit(args, f"[dry-run] NIM pada preview: {', '.join(row[nim_column] for row in preview[:10])}")
            if len(preview) > 10:
                _emit(args, f"[dry-run] ... dan {len(preview) - 10} NIM lainnya pada preview.")
            if args.limit and len(records) > args.limit:
                _emit(args, f"[dry-run] Batas per proses: {args.limit}; total roster: {len(records)}.")
            return 0

    base_url = config.base_url
    target_url = _ensure_same_origin(urljoin(base_url, args.url), base_url) if args.url else None
    if args.submit and not target_url:
        raise ValueError("Submit memerlukan --url halaman tambah mahasiswa yang aktif pada sesi eCampuz ini.")
    state_path = Path(args.state_file) if args.state_file else None
    log_path = Path(args.log_file) if args.log_file else None
    state: dict[str, Any] = {"fingerprint": "", "rows": {}}
    pending_records = records
    campaign_id = ""
    if args.submit and input_path is not None:
        target_parts = urlsplit(target_url)
        campaign_id = args.campaign_id or (
            parse_qs(target_parts.query).get("w", [""])[0]
            or target_parts.path + "?" + target_parts.query
        )
        fingerprint = _campaign_fingerprint(input_path, base_url, campaign_id)
        campaign_hash = hashlib.sha256(campaign_id.encode("utf-8")).hexdigest()[:8]
        state_path = state_path or STATE_DIR / f"yudisium_{input_path.stem}_{campaign_hash}_state.json"
        log_path = log_path or config.out_dir / f"yudisium_{input_path.stem}_{campaign_hash}_log.csv"
        state = _load_state(state_path, fingerprint)
        pending_records = _pending_records(
            records,
            nim_column,
            state,
            args.retry_failed,
        )
        pending_records = pending_records[: args.limit]
        if not pending_records:
            _emit(args, "[ok] Tidak ada baris tersisa pada batas yang diminta.")
            return 0
    username, password = _resolve_profile_credentials(
        config,
        allow_prompt=not getattr(args, "non_interactive", False),
    )
    profile_dir = Path(args.user_data_dir) if args.user_data_dir else STATE_DIR / "yudisium_browser_profile"
    profile_dir.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as playwright:
        launch_kwargs = {
            "user_data_dir": str(profile_dir.resolve()),
            "headless": False,
            "locale": "id-ID",
            "user_agent": config.user_agent,
            "viewport": {"width": 1366, "height": 900},
        }
        try:
            context = playwright.chromium.launch_persistent_context(**launch_kwargs)
        except Exception:
            context = playwright.chromium.launch_persistent_context(channel="chrome", **launch_kwargs)

        _inject_session_cookies(context, config)
        page = context.pages[0] if context.pages else context.new_page()
        try:
            if not page.url.startswith(base_url):
                page.goto(base_url, wait_until="domcontentloaded")
            pause_callback = getattr(args, "pause_for_manual_login", None)
            freshly_navigated_url = _manual_login(
                page,
                base_url,
                username,
                password,
                pause_callback=pause_callback,
            )
            registration_url = freshly_navigated_url or target_url
            if not registration_url and args.inspect:
                print("Arahkan browser ke halaman Tambah - Mahasiswa Lain yang ingin diperiksa.")
                input("Setelah halaman siap, tekan Enter di terminal ini: ")
                registration_url = _ensure_same_origin(page.url, base_url)
                if "/kelulusan/" not in urlsplit(registration_url).path.casefold():
                    raise RuntimeError("Halaman saat ini bukan halaman modul kelulusan eCampuz.")
            if not registration_url:
                raise ValueError("URL halaman yudisium tidak tersedia.")
            if page.url != registration_url:
                page.goto(registration_url, wait_until="domcontentloaded", timeout=args.timeout * 1000)
            try:
                page.wait_for_load_state("load", timeout=args.timeout * 1000)
            except Exception:
                pass
            if _is_login_page(page):
                raise RuntimeError("Halaman pendaftaran mengarah kembali ke login.")

            if not _is_registration_page(page):
                try:
                    discovered_url = _discover_registration_page(
                        page, base_url, args.timeout * 1000
                    )
                except Exception as exc:
                    discovered_url = None
                    discovery_error = exc
                else:
                    discovery_error = None
                if discovered_url:
                    registration_url = discovered_url
                    _emit(
                        args,
                        "[i] URL Pendaftaran Yudisium diperbarui dari menu eCampuz "
                        "yang sedang aktif.",
                    )
                elif not args.inspect:
                    reason = "URL halaman pendaftaran tidak dapat ditemukan."
                    if discovery_error:
                        reason = f"{reason} {discovery_error}"
                    raise RuntimeError(reason)

            if args.inspect:
                _print_inspection(page)
                return 0

            assert input_path is not None and log_path is not None and state_path is not None
            if args.submit:
                _ensure_registration_page(page)
                if campaign_id:
                    _select_semester(page, campaign_id)
            _emit(args, f"[siap] {len(pending_records)} pendaftaran akan dikirim berurutan.")
            _emit(args, f"[siap] Log: {log_path}")
            _emit(args, f"[siap] Profil browser tersimpan di: {profile_dir}")
            confirmation = f"DAFTAR YUDISIUM {len(pending_records)}"
            if not getattr(args, "confirmed", False):
                answer = input(f"Ketik {confirmation!r} untuk melanjutkan: ").strip()
                if answer != confirmation:
                    _emit(args, "[batal] Konfirmasi tidak cocok; tidak ada pendaftaran dikirim.")
                    return 1

            for index, record in enumerate(pending_records, start=1):
                if getattr(args, "cancel_check", lambda: False)():
                    _emit(args, "[batal] Proses pendaftaran dibatalkan atas permintaan operator.")
                    return 1
                nim = record[nim_column]
                if _is_login_page(page):
                    refreshed_url = _manual_login(
                        page,
                        base_url,
                        username,
                        password,
                        pause_callback=pause_callback,
                    )
                    registration_url = refreshed_url or registration_url
                    if page.url != registration_url:
                        page.goto(registration_url, wait_until="domcontentloaded", timeout=args.timeout * 1000)

                if page.url != registration_url:
                    page.goto(registration_url, wait_until="domcontentloaded", timeout=args.timeout * 1000)

                dialog_messages: list[str] = []
                confirmation_seen = False

                def accept_dialog(dialog) -> None:
                    nonlocal confirmation_seen
                    dialog_messages.append(dialog.message.strip())
                    if re.search(r"apakah|yakin|konfirmasi|are you sure", dialog.message, re.I):
                        confirmation_seen = True
                    dialog.accept()

                page.on("dialog", accept_dialog)
                search_error = None
                already_registered_during_search = False
                while True:
                    try:
                        if not args.no_reset_filter:
                            _try_reset_filter(page)
                        if campaign_id:
                            _select_semester(page, campaign_id)
                        nim_start, nim_end = _find_nim_range(page)
                        nim_start.fill("")
                        nim_end.fill("")
                        nim_start.fill(nim)
                        nim_end.fill(nim)
                        _button_by_text(page, "TAMPILKAN").click()
                        _wait_for_results(page, nim, args.timeout * 1000)
                        row = _find_exact_student_row(page, nim)
                        action = _student_action(row)
                        break
                    except Exception as exc:
                        if _is_login_page(page):
                            _emit(args, f"[jeda] Sesi habis saat mencari {nim}; login dan CAPTCHA manual, lalu bot akan mengulang NIM ini.")
                            refreshed_url = _manual_login(
                                page,
                                base_url,
                                username,
                                password,
                                pause_callback=pause_callback,
                            )
                            registration_url = refreshed_url or registration_url
                            if page.url != registration_url:
                                page.goto(registration_url, wait_until="domcontentloaded", timeout=args.timeout * 1000)
                            continue
                        # Mahasiswa yang sudah terdaftar sering tidak lagi
                        # muncul di daftar Tambah. Cek Pengelolaan sebelum
                        # menandainya gagal agar batch dapat diulang dengan aman.
                        if (
                            getattr(args, "verify_existing", True)
                            and campaign_id
                        ):
                            try:
                                already_registered = _management_has_registration(
                                    page,
                                    nim,
                                    campaign_id,
                                    args.timeout * 1000,
                                )
                            except Exception:
                                already_registered = None
                            if already_registered:
                                page.remove_listener("dialog", accept_dialog)
                                reason = "Sudah terdaftar di Pengelolaan Yudisium."
                                state.setdefault("rows", {})[nim] = {
                                    "status": "submitted",
                                    "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                                    "reason": reason,
                                }
                                _save_state(state_path, state)
                                _log_result(log_path, nim, "Sukses", reason)
                                _emit(
                                    args,
                                    f"[{index}/{len(pending_records)}] SUDAH ADA {nim}: {reason}",
                                )
                                already_registered_during_search = True
                                break
                        search_error = exc
                        break
                if already_registered_during_search:
                    continue
                if search_error is not None:
                    page.remove_listener("dialog", accept_dialog)
                    result = {
                        "status": "failed",
                        "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                        "reason": f"{search_error.__class__.__name__}: {search_error}",
                    }
                    state.setdefault("rows", {})[nim] = result
                    _save_state(state_path, state)
                    _log_result(log_path, nim, "Gagal", result["reason"])
                    _emit(args, f"[{index}/{len(pending_records)}] GAGAL {nim}: {search_error}")
                    continue

                if getattr(args, "verify_existing", True) and campaign_id:
                    try:
                        already_registered = _management_has_registration(
                            page,
                            nim,
                            campaign_id,
                            args.timeout * 1000,
                        )
                    except Exception as exc:
                        already_registered = None
                        verification_error = f"{exc.__class__.__name__}: {exc}"
                    else:
                        verification_error = ""

                    if already_registered is None:
                        page.remove_listener("dialog", accept_dialog)
                        reason = (
                            "Tidak dapat memverifikasi status mahasiswa di "
                            "Pengelolaan Yudisium."
                        )
                        if verification_error:
                            reason = f"{reason} {verification_error}"
                        state.setdefault("rows", {})[nim] = {
                            "status": "uncertain",
                            "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                            "reason": reason,
                        }
                        _save_state(state_path, state)
                        _log_result(log_path, nim, "Tidak pasti", reason)
                        artifact = _save_uncertain_artifact(page, nim)
                        if artifact is not None:
                            _emit(args, f"[bukti] Screenshot: {artifact.with_suffix('.png')}")
                            _emit(args, f"[bukti] HTML: {artifact.with_suffix('.html')}")
                        _emit(
                            args,
                            f"[{index}/{len(pending_records)}] HASIL TIDAK PASTI {nim}; "
                            f"verifikasi gagal: {reason}",
                        )
                        return 2

                    if already_registered:
                        page.remove_listener("dialog", accept_dialog)
                        reason = "Sudah terdaftar di Pengelolaan Yudisium."
                        state.setdefault("rows", {})[nim] = {
                            "status": "submitted",
                            "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                            "reason": reason,
                        }
                        _save_state(state_path, state)
                        _log_result(log_path, nim, "Sukses", reason)
                        _emit(
                            args,
                            f"[{index}/{len(pending_records)}] SUDAH ADA {nim}: {reason}",
                        )
                        continue

                dialog_messages.clear()
                body_before_action = page.locator("body").inner_text()
                body_before_lines = {line.strip() for line in body_before_action.splitlines() if line.strip()}
                body_before_result = body_before_action
                try:
                    if getattr(args, "cancel_check", lambda: False)():
                        page.remove_listener("dialog", accept_dialog)
                        _emit(args, f"[batal] Dibatalkan sebelum submit {nim}.")
                        return 1
                    state.setdefault("rows", {})[nim] = {
                        "status": "submitting",
                        "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                    }
                    _save_state(state_path, state)
                    action.click(timeout=args.timeout * 1000)
                    try:
                        modal_confirm = page.locator(".swal2-container .swal2-confirm, .modal.show button.btn-primary, .modal.show button:has-text('Ya')").first
                        if modal_confirm.is_visible(timeout=1000):
                            modal_confirm.click()
                    except Exception:
                        pass
                    try:
                        page.wait_for_load_state("domcontentloaded", timeout=args.timeout * 1000)
                    except Exception:
                        pass
                    try:
                        page.locator("input[name='mhsNiu']").nth(0).wait_for(
                            state="attached", timeout=min(args.timeout * 1000, 5000)
                        )
                    except Exception:
                        pass

                    detail_submit = _registration_detail_submit(page, nim)
                    if detail_submit is not None:
                        body_before_result = page.locator("body").inner_text()
                        body_before_lines = {
                            line.strip()
                            for line in body_before_result.splitlines()
                            if line.strip()
                        }
                        _fill_registration_date(
                            page, getattr(args, "registration_date", "") or ""
                        )
                        if getattr(args, "cancel_check", lambda: False)():
                            page.remove_listener("dialog", accept_dialog)
                            _emit(args, f"[batal] Dibatalkan sebelum submit detail {nim}.")
                            return 1
                        dialog_messages.clear()
                        confirmation_seen = False
                        detail_submit.click(timeout=args.timeout * 1000)
                        try:
                            modal_confirm = page.locator(".swal2-container .swal2-confirm, .modal.show button.btn-primary, .modal.show button:has-text('Ya')").first
                            if modal_confirm.is_visible(timeout=1000):
                                modal_confirm.click()
                        except Exception:
                            pass
                        try:
                            page.wait_for_function(
                                """oldText => {
                                  const body = document.body ? document.body.innerText : '';
                                  const notices = Array.from(document.querySelectorAll('.alert, .toast, [role="alert"], .swal2-html-container'));
                                  return body !== oldText || notices.some(node => (node.innerText || '').trim());
                                }""",
                                arg=body_before_result,
                                timeout=args.timeout * 1000,
                            )
                        except Exception:
                            try:
                                page.wait_for_load_state("domcontentloaded", timeout=args.timeout * 1000)
                            except Exception:
                                pass
                    elif confirmation_seen:
                        page.wait_for_function(
                            """oldText => {
                              const body = document.body ? document.body.innerText : '';
                              const notices = Array.from(document.querySelectorAll('.alert, .toast, [role="alert"], .swal2-html-container'));
                              return body !== oldText || notices.some(node => (node.innerText || '').trim());
                            }""",
                            arg=body_before_action,
                            timeout=args.timeout * 1000,
                        )
                    else:
                        try:
                            page.wait_for_load_state("domcontentloaded", timeout=args.timeout * 1000)
                        except Exception:
                            pass
                    _ensure_same_origin(page.url, base_url)
                except Exception as exc:
                    page.remove_listener("dialog", accept_dialog)
                    result = {
                        "status": "uncertain",
                        "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                        "reason": f"{exc.__class__.__name__}: {exc}",
                    }
                    state["rows"][nim] = result
                    _save_state(state_path, state)
                    _log_result(log_path, nim, "Tidak pasti", result["reason"])
                    artifact = _save_uncertain_artifact(page, nim)
                    if artifact is not None:
                        _emit(args, f"[bukti] Screenshot: {artifact.with_suffix('.png')}")
                        _emit(args, f"[bukti] HTML: {artifact.with_suffix('.html')}")
                    _emit(args, f"[{index}/{len(pending_records)}] HASIL TIDAK PASTI {nim}; dihentikan: {exc}")
                    return 2
                finally:
                    try:
                        page.remove_listener("dialog", accept_dialog)
                    except Exception:
                        pass

                result_status, reason = _classify_registration_result(
                    dialog_messages,
                    _registration_notices(page),
                    "\n".join(
                        line for line in page.locator("body").inner_text().splitlines()
                        if line.strip() and line.strip() not in body_before_lines
                    ),
                )
                if (
                    result_status in {"failed", "uncertain"}
                    and getattr(args, "verify_existing", True)
                    and campaign_id
                ):
                    try:
                        verified_after_submit = _management_has_registration(
                            page,
                            nim,
                            campaign_id,
                            args.timeout * 1000,
                        )
                    except Exception:
                        verified_after_submit = None
                    if verified_after_submit:
                        result_status = "submitted"
                        reason = "Terverifikasi ada di Pengelolaan Yudisium."
                    elif verified_after_submit is False and result_status == "uncertain":
                        reason = (
                            f"{reason} (tidak ditemukan di Pengelolaan Yudisium)"
                        )
                status = result_status
                state["rows"][nim] = {
                    "status": status,
                    "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
                    "reason": reason,
                }
                _save_state(state_path, state)
                display_status = {"submitted": "Sukses", "failed": "Gagal", "uncertain": "Tidak pasti"}[status]
                _log_result(log_path, nim, display_status, reason)
                _emit(args, f"[{index}/{len(pending_records)}] {display_status.upper()} {nim}: {reason}")
                if status == "uncertain":
                    artifact = _save_uncertain_artifact(page, nim)
                    if artifact is not None:
                        _emit(args, f"[bukti] Screenshot: {artifact.with_suffix('.png')}")
                        _emit(args, f"[bukti] HTML: {artifact.with_suffix('.html')}")
                    return 2
                if index < len(pending_records):
                    total_sleep = max(args.delay, 1.5)
                    steps = int(total_sleep / 0.2)
                    for _ in range(max(1, steps)):
                        if getattr(args, "cancel_check", lambda: False)():
                            _emit(args, "[batal] Proses pendaftaran dibatalkan saat jeda.")
                            return 1
                        time.sleep(0.2)
            return 0
        finally:
            context.close()